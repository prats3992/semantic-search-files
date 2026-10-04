import docx
import pytest

from semantic_search.content_extraction.extractor import ContentExtractor


@pytest.fixture
def gemini_calls(monkeypatch):
    """Records Gemini uploads instead of calling the real API."""
    calls = []

    def fake_gemini(self, file_path):
        calls.append(file_path)
        return "text from gemini"

    monkeypatch.setattr(ContentExtractor, "_extract_text_with_gemini", fake_gemini)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    return calls


def test_reads_plain_text_and_source_files(tmp_path):
    for name in ("notes.txt", "README.md", "script.py"):
        (tmp_path / name).write_text(f"contents of {name}")

    extractor = ContentExtractor()

    for name in ("notes.txt", "README.md", "script.py"):
        assert extractor.extract_text(str(tmp_path / name)) == f"contents of {name}"


def test_extension_matching_is_case_insensitive(tmp_path):
    path = tmp_path / "LOUD.TXT"
    path.write_text("hello")

    assert ContentExtractor().extract_text(str(path)) == "hello"


def test_extracts_docx_paragraphs(tmp_path):
    path = tmp_path / "report.docx"
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.save(path)

    assert ContentExtractor().extract_text(str(path)) == "First paragraph.\nSecond paragraph."


def test_corrupt_docx_returns_none(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a docx")

    assert ContentExtractor().extract_text(str(path)) is None


def test_pdf_is_skipped_without_api_key(tmp_path, gemini_calls):
    path = tmp_path / "paper.pdf"
    path.write_bytes(b"%PDF-1.4")

    assert ContentExtractor().extract_text(str(path)) is None
    assert gemini_calls == []


@pytest.mark.parametrize("name", ["paper.pdf", "scan.PNG", "photo.jpg", "photo.jpeg", "diagram.webp"])
def test_pdfs_and_images_use_gemini_when_key_is_set(tmp_path, gemini_calls, name):
    path = tmp_path / name
    path.write_bytes(b"binary")

    assert ContentExtractor(gemini_api_key="test-key").extract_text(str(path)) == "text from gemini"
    assert gemini_calls == [str(path)]


def test_api_key_is_read_from_environment(tmp_path, gemini_calls, monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "env-key")
    path = tmp_path / "paper.pdf"
    path.write_bytes(b"%PDF-1.4")

    assert ContentExtractor().extract_text(str(path)) == "text from gemini"


@pytest.mark.parametrize("name", ["archive.zip", "program.exe", "data.bin", "no_extension"])
def test_other_binaries_are_never_sent_to_gemini(tmp_path, gemini_calls, name):
    path = tmp_path / name
    path.write_bytes(b"\x00\x01")

    assert ContentExtractor(gemini_api_key="test-key").extract_text(str(path)) is None
    assert gemini_calls == []
